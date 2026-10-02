"""Bin-label and missing-label reports for a 5S audit, published as an Excel workbook.

Every one of the 4,000 bins either has a physical label record or does not;
LABEL_MISSING_RATE seeds a non-trivial fraction of bins with no label so the
missing-label report is not vacuously empty. The workbook has three sheets:
one row per bin with its label status (full audit trail), a filtered sheet
of only the bins missing a label (the actionable 5S worklist), and a summary.
"""
from __future__ import annotations

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ssk import reference as ref


def generate_label_status(bins: pd.DataFrame, rng: np.random.Generator | None = None) -> pd.DataFrame:
    """One row per bin: whether a physical label record exists for it."""
    if rng is None:
        rng = np.random.default_rng(ref.RNG_SEED_AUDIT + 1)

    n = len(bins)
    missing = rng.random(n) < ref.LABEL_MISSING_RATE
    df = bins[["bin_id", "aisle", "bay", "shelf"]].copy()
    df["label_status"] = np.where(missing, "MISSING", "PRESENT")
    df["location"] = (
        "A" + df["aisle"].astype(str) + "-B" + df["bay"].astype(str) + "-S" + df["shelf"].astype(str)
    )
    return df


def write_5s_audit_workbook(label_status: pd.DataFrame, path: str) -> None:
    """Write the three-sheet 5S audit workbook: Bin Labels, Missing Labels, Summary."""
    wb = Workbook()

    ws_all = wb.active
    ws_all.title = "Bin Labels"
    header_font = Font(bold=True)
    missing_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    cols = ["bin_id", "location", "aisle", "bay", "shelf", "label_status"]
    ws_all.append(cols)
    for cell in ws_all[1]:
        cell.font = header_font
    for row in label_status[cols].itertuples(index=False):
        ws_all.append(list(row))
    for row_idx in range(2, ws_all.max_row + 1):
        status_cell = ws_all.cell(row=row_idx, column=cols.index("label_status") + 1)
        if status_cell.value == "MISSING":
            for c in range(1, len(cols) + 1):
                ws_all.cell(row=row_idx, column=c).fill = missing_fill

    ws_missing = wb.create_sheet("Missing Labels")
    ws_missing.append(cols)
    for cell in ws_missing[1]:
        cell.font = header_font
    missing_df = label_status[label_status["label_status"] == "MISSING"][cols]
    for row in missing_df.itertuples(index=False):
        ws_missing.append(list(row))

    ws_summary = wb.create_sheet("Summary")
    total = len(label_status)
    n_missing = int((label_status["label_status"] == "MISSING").sum())
    n_present = total - n_missing
    ws_summary.append(["metric", "value"])
    for cell in ws_summary[1]:
        cell.font = header_font
    ws_summary.append(["total_bins", total])
    ws_summary.append(["labels_present", n_present])
    ws_summary.append(["labels_missing", n_missing])
    ws_summary.append(["missing_pct", round(100.0 * n_missing / total, 4) if total else 0.0])

    wb.save(path)
