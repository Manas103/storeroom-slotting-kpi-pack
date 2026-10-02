"""Test suite: bin-capacity invariant for both assignments, the oracle-vs-
heuristic never-shorter-than-optimal property, a hand-verifiable oracle case,
KPI lineage coverage, and the 5S audit workbook's row accounting."""
from __future__ import annotations

import os
import sys

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from ssk import reference as ref
from ssk.layout import build_bins, coords_lookup, bin_distance, entrance_distance
from ssk.demand import generate_materials, generate_picks, assign_pick_lists, material_pick_frequency
from ssk.slotting import compute_required_units, baseline_assignment, reslotted_assignment, assert_capacity_respected
from ssk.oracle import optimal_route_distance
from ssk.travel import nearest_neighbor_route_distance
from ssk.publish import build_storeroom, build_demand, build_slotting, sample_pick_lists_for_oracle, travel_benchmark
from ssk.kpis import simulate_inventory, generate_cycle_counts, connect, all_kpis, lineage_coverage
from ssk.audit import generate_label_status, write_5s_audit_workbook


@pytest.fixture(scope="module")
def small_world():
    """Full-scale materials/picks generation (required: pick_weight is a
    Zipf distribution over all N_MATERIALS and must sum to 1), but a fresh,
    distinct seed set from the benchmark so this is an independent run, not
    a replay of scripts/run_benchmark.py's exact numbers."""
    rng_bins = np.random.default_rng(11)
    bins = build_bins(rng_bins)
    materials = generate_materials(np.random.default_rng(12))
    picks = generate_picks(materials, np.random.default_rng(13))
    picks = assign_pick_lists(picks, np.random.default_rng(14))
    pick_freq = material_pick_frequency(picks)
    materials_req = compute_required_units(materials, pick_freq)
    return bins, materials, picks, materials_req


def test_bin_count_and_coords_match_reference():
    bins = build_bins(np.random.default_rng(1))
    assert len(bins) == ref.N_BINS == 4000
    coords = coords_lookup(bins)
    assert len(coords) == ref.N_BINS
    assert all(len(c) == 3 for c in coords.values())


def test_baseline_and_reslotted_respect_bin_capacity(small_world):
    bins, materials, picks, materials_req = small_world
    baseline = baseline_assignment(bins, materials_req)
    reslotted = reslotted_assignment(bins, materials_req)
    assert_capacity_respected(baseline, bins, materials_req)
    assert_capacity_respected(reslotted, bins, materials_req)
    # Every material placed exactly once under each assignment.
    assert set(baseline.keys()) == set(materials_req["material_id"])
    assert set(reslotted.keys()) == set(materials_req["material_id"])


def test_reslotted_puts_high_frequency_materials_closer_to_entrance(small_world):
    bins, materials, picks, materials_req = small_world
    reslotted = reslotted_assignment(bins, materials_req)
    dist = dict(zip(bins["bin_id"], bins["dist_from_entrance"]))
    merged = materials_req.copy()
    merged["bin_id"] = merged["material_id"].map(reslotted)
    merged["dist"] = merged["bin_id"].map(dist)
    top_decile = merged.nlargest(max(1, len(merged) // 10), "pick_count")
    bottom_decile = merged.nsmallest(max(1, len(merged) // 10), "pick_count")
    assert top_decile["dist"].mean() <= bottom_decile["dist"].mean()


def test_oracle_hand_verifiable_two_bin_case():
    """A hand-computable case: two bins on the same aisle/bay, different
    shelves. The optimal round trip must visit both in increasing distance
    order from the entrance and its length is checkable by hand."""
    coords = [(0, 0, 0), (0, 0, 5)]
    dist = optimal_route_distance(coords)
    expected = entrance_distance((0, 0, 0)) + bin_distance((0, 0, 0), (0, 0, 5)) + entrance_distance((0, 0, 5))
    assert abs(dist - expected) < 1e-9


def test_oracle_raises_above_its_bounded_size():
    coords = [(0, i, 0) for i in range(ref.PICK_LIST_MAX_SIZE + 1)]
    with pytest.raises(ValueError):
        optimal_route_distance(coords)


def test_heuristic_never_beats_the_oracle(small_world):
    """The fast nearest-neighbor heuristic must never report a route shorter
    than the oracle's true optimum for the same stops, checked over many
    random small pick lists, not just the one case travel_benchmark samples."""
    rng = np.random.default_rng(99)
    bins, *_ = small_world
    coords_pool = list(coords_lookup(bins).values())
    for _ in range(200):
        n = rng.integers(2, ref.PICK_LIST_MAX_SIZE + 1)
        sample = [coords_pool[i] for i in rng.choice(len(coords_pool), size=n, replace=False)]
        oracle_dist = optimal_route_distance(sample)
        heuristic_dist = nearest_neighbor_route_distance(sample)
        assert heuristic_dist >= oracle_dist - 1e-6


def test_oracle_sample_lists_are_within_bounds():
    bins = build_storeroom()
    materials, picks = build_demand()
    list_ids = sample_pick_lists_for_oracle(picks)
    sizes = picks[picks["list_id"].isin(list_ids)].groupby("list_id").size()
    assert sizes.between(2, ref.PICK_LIST_MAX_SIZE).all()
    assert len(list_ids) <= ref.ORACLE_SAMPLE_SIZE


def test_kpi_lineage_coverage_is_complete(small_world):
    bins, materials, picks, materials_req = small_world
    inventory_daily = simulate_inventory(materials_req, picks)
    cycle_counts = generate_cycle_counts(inventory_daily)
    con = connect(bins, materials_req, picks, inventory_daily, cycle_counts)
    results = all_kpis(con)
    assert len(results) == 6
    coverage = lineage_coverage(results)
    assert coverage == 100.0
    for r in results:
        assert r.lineage_row_count > 0
        assert "SELECT" in r.lineage_query.upper()


def test_5s_audit_workbook_accounts_for_every_bin(tmp_path):
    bins = build_bins(np.random.default_rng(1))
    label_status = generate_label_status(bins)
    assert len(label_status) == ref.N_BINS
    n_missing = int((label_status["label_status"] == "MISSING").sum())
    assert n_missing > 0  # the missing-label report must not be vacuously empty

    out_path = str(tmp_path / "5s_audit_test.xlsx")
    write_5s_audit_workbook(label_status, out_path)
    assert os.path.exists(out_path)

    import openpyxl
    wb = openpyxl.load_workbook(out_path)
    assert set(wb.sheetnames) == {"Bin Labels", "Missing Labels", "Summary"}
    ws_all = wb["Bin Labels"]
    assert ws_all.max_row - 1 == ref.N_BINS  # header + one row per bin
    ws_missing = wb["Missing Labels"]
    assert ws_missing.max_row - 1 == n_missing
